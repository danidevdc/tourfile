
from flask import Flask, request, send_file, jsonify
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from datetime import datetime
import io
import traceback

app = Flask(__name__)

def apply_borders(ws, min_row, max_row, min_col, max_col):
    thin_border = Border(
        left=Side(style='thin'), 
        right=Side(style='thin'),
        top=Side(style='thin'), 
        bottom=Side(style='thin')
    )
    for row_idx in range(min_row, max_row + 1):
        for col_idx in range(min_col, max_col + 1):
            ws.cell(row=row_idx, column=col_idx).border = thin_border

def generate_excel_in_memory(data):
    try:
        wb = Workbook()
        ws = wb.active
        ws.title = "CajaChica"

        file_number = data.get("fileNumber", "N/A")
        guide_name = data.get("guideName", "N/A").upper()
        group_name = data.get("groupName", "N/A")
        pax_count = int(data.get("paxCount", 0))
        expense_items = data.get("expenseItems", [])

        header_font = Font(name='Calibri', size=11, bold=True)
        
        center_alignment = Alignment(horizontal="center", vertical="center")
        right_alignment = Alignment(horizontal="right", vertical="center")
        left_alignment = Alignment(horizontal="left", vertical="center")
        
        currency_format = '#,##0.00'
        date_format_ddmmyy = 'dd/mm/yy'

        ws['A1'] = "CAJA CHICA GUIA"
        ws['A1'].font = Font(name='Calibri', size=14, bold=True)
        ws['A1'].alignment = left_alignment
        ws.merge_cells('A1:G1')

        ws['A2'] = "FILE:"
        ws['A2'].font = header_font
        ws['B2'] = file_number
        ws['B2'].alignment = center_alignment
        ws.merge_cells('B2:C2')

        ws['D2'] = "NOMBRE GUIA:"
        ws['D2'].font = header_font
        ws['E2'] = guide_name
        ws['E2'].alignment = center_alignment
        ws.merge_cells('E2:F2')

        ws['A3'] = "NOMBRE Y Nº DE PAX:"
        ws['A3'].font = header_font
        ws.merge_cells('A3:C3')

        ws['D3'] = group_name
        ws['D3'].alignment = center_alignment
        ws.merge_cells('D3:E3')

        ws['F3'] = "Nº"
        ws['F3'].font = header_font
        ws['G3'] = pax_count
        ws['G3'].font = header_font
        ws['G3'].alignment = right_alignment

        headers = ["FECHA", "CANT", "DETALLE DEL GASTO", "", "P. UNIT", "TOTAL Bs.", "VoB OPS"]
        for col_num, header_text in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col_num, value=header_text)
            cell.font = header_font
            cell.alignment = left_alignment
        ws.merge_cells('C4:D4')

        current_row = 5
        for item in expense_items:
            date_cell = ws.cell(row=current_row, column=1)
            date_val_str = item.get("date", "")
            if date_val_str:
                try:
                    date_obj = datetime.strptime(date_val_str, "%d/%m/%y")
                    date_cell.value = date_obj
                    date_cell.number_format = date_format_ddmmyy
                except (ValueError, TypeError):
                     # Handle cases where date is not in the expected format or is already a date object.
                     # This can be adjusted based on expected data formats.
                    date_cell.value = date_val_str
            else:
                date_cell.value = ""
            date_cell.alignment = center_alignment
            
            quantity_str = str(item.get("quantity", "1"))
            if quantity_str.startswith("="):
                ws.cell(row=current_row, column=2, value=quantity_str).alignment = center_alignment
            else:
                try:
                    ws.cell(row=current_row, column=2, value=float(quantity_str)).alignment = center_alignment
                except (ValueError, TypeError):
                    ws.cell(row=current_row, column=2, value=quantity_str).alignment = center_alignment

            ws.cell(row=current_row, column=3, value=item.get("detail", "")).alignment = center_alignment
            ws.merge_cells(start_row=current_row, start_column=3, end_row=current_row, end_column=4)
            
            unit_price_val = item.get("unitPrice", 0)
            try:
                unit_price_val = float(unit_price_val)
            except (ValueError, TypeError):
                unit_price_val = 0
            
            unit_price_cell = ws.cell(row=current_row, column=5, value=unit_price_val)
            unit_price_cell.number_format = currency_format
            unit_price_cell.alignment = right_alignment

            total_formula = f"=B{current_row}*E{current_row}"
            total_cell = ws.cell(row=current_row, column=6, value=total_formula)
            total_cell.number_format = currency_format
            total_cell.alignment = right_alignment
            
            ws.cell(row=current_row, column=7, value=item.get("vobOps", "")).alignment = center_alignment
            current_row += 1

        ws.cell(row=current_row, column=3, value="GASTO TOTAL").font = header_font
        ws.cell(row=current_row, column=3).alignment = center_alignment
        ws.merge_cells(start_row=current_row, start_column=3, end_row=current_row, end_column=5)

        grand_total_formula = f"=SUM(F5:F{current_row-1})" if expense_items else "0"
        grand_total_cell = ws.cell(row=current_row, column=6, value=grand_total_formula)
        grand_total_cell.font = header_font
        grand_total_cell.number_format = currency_format
        grand_total_cell.alignment = right_alignment
        
        ws.column_dimensions['A'].width = 12
        ws.column_dimensions['B'].width = 9.14
        ws.column_dimensions['C'].width = 9.14
        ws.column_dimensions['D'].width = 14.85
        ws.column_dimensions['E'].width = 9.14
        ws.column_dimensions['F'].width = 9.14
        ws.column_dimensions['G'].width = 9.14

        max_data_row = current_row
        apply_borders(ws, min_row=1, max_row=max_data_row, min_col=1, max_col=7)

        in_memory_fp = io.BytesIO()
        wb.save(in_memory_fp)
        in_memory_fp.seek(0)
        return in_memory_fp

    except Exception:
        traceback.print_exc()
        raise

@app.route("/", methods=['POST'])
def handle_excel_generation():
    if not request.is_json:
        return jsonify({"error": "Request must be JSON"}), 400

    report_data = request.get_json()
    if not report_data:
        return jsonify({"error": "Bad Request: Missing report data"}), 400

    try:
        excel_buffer = generate_excel_in_memory(report_data)

        safe_group_name = str(report_data.get("groupName", "grupo")).replace('/', '_').replace('\\', '_')
        safe_guide_name = str(report_data.get("guideName", "guia")).replace('/', '_').replace('\\', '_')
        safe_file_number = str(report_data.get("fileNumber", "file")).replace('/', '_').replace('\\', '_')
        start_date = report_data.get("startDate", "")
        
        final_constructed_file_name = f"G.O. {start_date} - {safe_group_name} - {safe_guide_name} - {safe_file_number}.xlsx"

        return send_file(
            excel_buffer,
            as_attachment=True,
            download_name=final_constructed_file_name,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )

    except Exception as e:
        print(f"Internal server error: {e}")
        traceback.print_exc()
        return jsonify({"error": "An internal error occurred during Excel generation."}), 500

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=8080)
