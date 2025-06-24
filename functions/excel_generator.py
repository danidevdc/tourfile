#!/usr/bin/env python3

import sys
import json
import os
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def load_data(file_path):
    """Carga los datos desde el archivo JSON"""
    try:
        with open(file_path, 'r', encoding='utf-8') as file:
            return json.load(file)
    except Exception as e:
        print(f"Error cargando datos: {e}")
        return None

def create_workbook(data):
    """Crea el workbook de Excel con los datos"""
    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte de Cajas Chicas"
    
    # Estilos
    header_font = Font(name='Arial', size=12, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
    cell_font = Font(name='Arial', size=10)
    
    border_style = Side(border_style='thin', color='000000')
    border = Border(left=border_style, right=border_style, top=border_style, bottom=border_style)
    
    center_alignment = Alignment(horizontal='center', vertical='center')
    left_alignment = Alignment(horizontal='left', vertical='center')
    right_alignment = Alignment(horizontal='right', vertical='center')
    
    # Headers
    headers = [
        'Fecha', 'Descripción', 'Categoría', 'Monto (Bs)', 
        'Responsable', 'Comprobante', 'Estado', 'Observaciones'
    ]
    
    # Escribir headers
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_alignment
        cell.border = border
    
    # Procesar datos
    current_row = 2
    total_amount = 0
    
    # Si hay elementos en el data
    if 'elements' in data and data['elements']:
        for element in data['elements']:
            try:
                # Mapear los datos del formulario
                fecha = element.get('fecha', '')
                descripcion = element.get('descripcion', '')
                categoria = element.get('categoria', '')
                monto = float(element.get('monto', 0))
                responsable = element.get('responsable', '')
                comprobante = element.get('comprobante', '')
                estado = element.get('estado', 'Pendiente')
                observaciones = element.get('observaciones', '')
                
                # Escribir fila
                row_data = [fecha, descripcion, categoria, monto, responsable, comprobante, estado, observaciones]
                
                for col_num, value in enumerate(row_data, 1):
                    cell = ws.cell(row=current_row, column=col_num)
                    cell.value = value
                    cell.font = cell_font
                    cell.border = border
                    
                    # Alineación específica por columna
                    if col_num == 4:  # Monto
                        cell.alignment = right_alignment
                        cell.number_format = '#,##0.00'
                    elif col_num in [1, 6]:  # Fecha y Comprobante
                        cell.alignment = center_alignment
                    else:
                        cell.alignment = left_alignment
                
                total_amount += monto
                current_row += 1
                
            except Exception as e:
                print(f"Error procesando elemento: {e}")
                continue
    
    # Agregar fila de total
    if current_row > 2:  # Si hay datos
        total_row = current_row + 1
        
        # Etiqueta "TOTAL"
        total_label_cell = ws.cell(row=total_row, column=3)
        total_label_cell.value = "TOTAL:"
        total_label_cell.font = Font(name='Arial', size=12, bold=True)
        total_label_cell.alignment = right_alignment
        total_label_cell.border = border
        
        # Valor total
        total_value_cell = ws.cell(row=total_row, column=4)
        total_value_cell.value = total_amount
        total_value_cell.font = Font(name='Arial', size=12, bold=True)
        total_value_cell.alignment = right_alignment
        total_value_cell.number_format = '#,##0.00'
        total_value_cell.border = border
        total_value_cell.fill = PatternFill(start_color='E6E6E6', end_color='E6E6E6', fill_type='solid')
    
    # Ajustar anchos de columna
    column_widths = {
        'A': 12,  # Fecha
        'B': 30,  # Descripción
        'C': 15,  # Categoría
        'D': 15,  # Monto
        'E': 20,  # Responsable
        'F': 15,  # Comprobante
        'G': 12,  # Estado
        'H': 25   # Observaciones
    }
    
    for col_letter, width in column_widths.items():
        ws.column_dimensions[col_letter].width = width
    
    # Agregar información adicional
    info_row = current_row + 3
    
    # Fecha de generación
    gen_date_cell = ws.cell(row=info_row, column=1)
    gen_date_cell.value = f"Generado el: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}"
    gen_date_cell.font = Font(name='Arial', size=9, italic=True)
    
    # Información del usuario si está disponible
    if 'user_info' in data:
        user_info = data['user_info']
        user_cell = ws.cell(row=info_row + 1, column=1)
        user_cell.value = f"Usuario: {user_info.get('name', 'N/A')}"
        user_cell.font = Font(name='Arial', size=9, italic=True)
    
    return wb

def main():
    if len(sys.argv) != 3:
        print("Uso: python excel_generator_cli.py <archivo_datos.json> <archivo_salida.xlsx>")
        sys.exit(1)
    
    input_file = sys.argv[1]
    output_file = sys.argv[2]
    
    print(f"📥 Cargando datos desde: {input_file}")
    print(f"📤 Archivo de salida: {output_file}")
    
    # Verificar que el archivo de entrada existe
    if not os.path.exists(input_file):
        print(f"❌ Error: No se encuentra el archivo {input_file}")
        sys.exit(1)
    
    # Cargar datos
    data = load_data(input_file)
    if data is None:
        print("❌ Error: No se pudieron cargar los datos")
        sys.exit(1)
    
    print(f"📊 Datos cargados correctamente")
    print(f"📊 Elementos a procesar: {len(data.get('elements', []))}")
    
    try:
        # Crear workbook
        wb = create_workbook(data)
        
        # Guardar archivo
        wb.save(output_file)
        print(f"✅ Archivo Excel creado exitosamente: {output_file}")
        
        # Verificar que el archivo se creó
        if os.path.exists(output_file):
            file_size = os.path.getsize(output_file)
            print(f"📏 Tamaño del archivo: {file_size} bytes")
        else:
            print("❌ Error: El archivo no se creó correctamente")
            sys.exit(1)
            
    except Exception as e:
        print(f"❌ Error creando Excel: {e}")
        import traceback
        traceback.print_exc()
        sys.exit(1)

if __name__ == "__main__":
    main()