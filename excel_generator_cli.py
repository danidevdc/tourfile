#!/usr/bin/env python3

import sys
import json
import os
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def load_data(file_path):
    """Carga los datos desde el archivo JSON"""
    try:
        with open(file_path, 'r', encoding='utf-8') as file:
            return json.load(file)
    except Exception as e:
        # Using stderr for error logging
        print(f"Error loading data: {e}", file=sys.stderr)
        return None

def create_workbook(data):
    """Crea el workbook de Excel con los datos"""
    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte de Cajas Chicas"
    
    # Estilos
    header_font = Font(name='Arial', size=12, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
    cell_font = Font(name='Arial', size=10)
    
    border_style = Side(border_style='thin', color='000000')
    border = Border(left=border_style, right=border_style, top=border_style, bottom=border_style)
    
    center_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left_alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    right_alignment = Alignment(horizontal='right', vertical='center')
    
    # Headers
    headers = [
        'Fecha', 'Descripción', 'Categoría', 'Monto (Bs)', 
        'Responsable', 'Comprobante', 'Estado', 'Observaciones'
    ]
    
    # Escribir headers
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_alignment
        cell.border = border
    
    # Procesar datos
    current_row = 2
    total_amount = 0
    
    if 'expenseItems' in data and data['expenseItems']:
        for element in data['expenseItems']:
            try:
                # Mapear los datos del formulario
                fecha = element.get('date', '')
                descripcion = element.get('detail', '')
                categoria = '' # Categoría no viene en los datos, se deja en blanco
                monto = float(element.get('total', 0))
                responsable = data.get('guideName', '')
                comprobante = '' # Comprobante no viene en los datos
                estado = 'Pendiente' # Estado por defecto
                observaciones = '' # Observaciones no vienen en los datos

                row_data = [fecha, descripcion, categoria, monto, responsable, comprobante, estado, observaciones]
                
                for col_num, value in enumerate(row_data, 1):
                    cell = ws.cell(row=current_row, column=col_num)
                    cell.value = value
                    cell.font = cell_font
                    cell.border = border
                    
                    if col_num == 4:
                        cell.alignment = right_alignment
                        cell.number_format = '#,##0.00'
                    elif col_num == 1:
                        cell.alignment = center_alignment
                    else:
                        cell.alignment = left_alignment
                
                total_amount += monto
                current_row += 1
                
            except Exception as e:
                print(f"Error processing element: {element}. Error: {e}", file=sys.stderr)
                continue
    
    # Fila de total
    if current_row > 2:
        total_row = current_row + 1
        ws.cell(row=total_row, column=3).value = "TOTAL:"
        ws.cell(row=total_row, column=3).font = Font(name='Arial', size=12, bold=True)
        ws.cell(row=total_row, column=3).alignment = right_alignment
        
        total_value_cell = ws.cell(row=total_row, column=4)
        total_value_cell.value = total_amount
        total_value_cell.font = Font(name='Arial', size=12, bold=True)
        total_value_cell.alignment = right_alignment
        total_value_cell.number_format = '#,##0.00'
        total_value_cell.fill = PatternFill(start_color='E6E6E6', end_color='E6E6E6', fill_type='solid')

    # Ajustar anchos de columna
    column_widths = {'A': 12, 'B': 45, 'C': 15, 'D': 15, 'E': 20, 'F': 15, 'G': 12, 'H': 25}
    for col_letter, width in column_widths.items():
        ws.column_dimensions[col_letter].width = width
    
    return wb

def main():
    if len(sys.argv) != 3:
        print("Usage: python excel_generator_cli.py <input_json_path> <output_xlsx_path>", file=sys.stderr)
        sys.exit(1)
    
    input_file = sys.argv[1]
    output_file = sys.argv[2]
    
    if not os.path.exists(input_file):
        print(f"Error: Input file not found at {input_file}", file=sys.stderr)
        sys.exit(1)
    
    data = load_data(input_file)
    if data is None:
        print("Error: Could not load data from JSON file.", file=sys.stderr)
        sys.exit(1)
    
    try:
        wb = create_workbook(data)
        wb.save(output_file)
    except Exception as e:
        print(f"Error creating Excel workbook: {e}", file=sys.stderr)
        import traceback
        traceback.print_exc(file=sys.stderr)
        sys.exit(1)

if __name__ == "__main__":
    main()
